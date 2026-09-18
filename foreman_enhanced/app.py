import io, os, json
import pandas as pd
import pypdf
from docx import Document
import streamlit as st
from dotenv import load_dotenv
import sys

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
SRC_DIR = os.path.join(CURRENT_DIR, "src")

if SRC_DIR not in sys.path:
    sys.path.insert(0, SRC_DIR)

from scrummaster_businessanalyst_multiagent.crew import ForemanBacklogCrew
from scrummaster_businessanalyst_multiagent.models import BacklogPackage
from scrummaster_businessanalyst_multiagent.services.dataset_loader import load_workbook
from scrummaster_businessanalyst_multiagent.services.board_health_service import run_health_checks
from scrummaster_businessanalyst_multiagent.services.sprint_planner import plan_sprints
from scrummaster_businessanalyst_multiagent.services.change_service import analyze_change
from scrummaster_businessanalyst_multiagent.tools.jira_tool import publish_backlog

load_dotenv()
st.set_page_config(page_title="Foreman AI Scrum Master", page_icon="⚡", layout="wide")

for key, default in {"package": None, "data": None, "health": None, "plan": None, "approved": False}.items():
    if key not in st.session_state:
        st.session_state[key] = default

def extract(file):
    ext = os.path.splitext(file.name)[1].lower()
    if ext == ".txt": return file.getvalue().decode("utf-8")
    if ext == ".pdf": return "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(file).pages)
    if ext == ".docx": return "\n".join(p.text for p in Document(file).paragraphs)
    return ""

def package_dict():
    p = st.session_state.package
    return p.model_dump() if hasattr(p, "model_dump") else p

st.title("⚡ Foreman AI Scrum Master")
st.caption("Traceable backlog shaping, dependency-safe sprint planning, board health and human-controlled publishing")
page = st.sidebar.radio("Journey", ["1. Intake", "2. Clarification", "3. Backlog", "4. Sprint Plan", "5. Change Request", "6. Board Health", "7. Publish"])

if page == "1. Intake":
    st.header("Intake")
    uploaded = st.file_uploader("Upload TXT, PDF or DOCX", type=["txt", "pdf", "docx"])
    text = st.text_area("Or paste raw demand", height=240)
    raw = extract(uploaded) if uploaded else text
    if st.button("Generate backlog", type="primary"):
        if not raw.strip(): st.error("Provide input first.")
        else:
            with st.spinner("Foreman is shaping the demand..."):
                output = ForemanBacklogCrew().crew().kickoff(inputs={"transcript": raw})
                if getattr(output, "pydantic", None): st.session_state.package = output.pydantic
                else: st.session_state.package = BacklogPackage.model_validate_json(output.raw)
            st.success("Backlog generated. Continue to Clarification.")
    st.divider()
    wb = st.file_uploader("Upload Foreman synthetic workbook", type=["xlsx"], key="workbook")
    if wb and st.button("Load workbook"):
        st.session_state.data = load_workbook(wb)
        st.success("Workbook loaded. Board Health and Sprint Plan are ready.")

elif page == "2. Clarification":
    st.header("Clarification")
    if not st.session_state.package: st.info("Generate a backlog first.")
    else:
        p = package_dict()
        st.write(p.get("normalized_demand", ""))
        answers = {}
        for i, q in enumerate(p.get("clarification_questions", [])):
            answers[q] = st.text_input(q, key=f"q{i}")
        if st.button("Apply answers"):
            for item in p["items"]:
                if item.get("missing_details") and all(answers.values()):
                    item["description"] += "\n\nHuman clarification:\n" + "\n".join(f"{q}: {a}" for q,a in answers.items())
                    item["missing_details"] = []
                    item["readiness_status"] = "READY"
            st.session_state.package = p
            st.success("Answers applied. Review the backlog before approval.")

elif page == "3. Backlog":
    st.header("Backlog")
    if not st.session_state.package: st.info("Generate a backlog first.")
    else:
        p = package_dict()
        for i, item in enumerate(p["items"]):
            with st.expander(f"{item['local_id']} | {item['issue_type']} | {item['summary']}", expanded=item.get("readiness_status") != "READY"):
                item["summary"] = st.text_input("Summary", item["summary"], key=f"s{i}")
                item["description"] = st.text_area("Description", item["description"], key=f"d{i}")
                c1,c2,c3 = st.columns(3)
                item["parent_local_id"] = c1.text_input("Parent local ID", item.get("parent_local_id") or "", key=f"p{i}") or None
                item["sad_section_id"] = c2.text_input("S-AD section", item.get("sad_section_id") or "", key=f"sad{i}") or None
                if item["issue_type"] in ["Story","Task","Bug"]:
                    item["story_points"] = c3.selectbox("Points", [2,3,5,8,13], index=[2,3,5,8,13].index(item.get("story_points")) if item.get("story_points") in [2,3,5,8,13] else 2, key=f"pts{i}")
                st.caption(f"Readiness: {item.get('readiness_status')} | Layer: {item.get('architecture_layer')} | Confidence: {item.get('estimation_confidence')}")
        st.session_state.package = p
        st.session_state.approved = st.checkbox("I approve this internal backlog for planning", value=st.session_state.approved)
        st.download_button("Download backlog JSON", json.dumps(p, indent=2), "foreman_backlog.json", "application/json")

elif page == "4. Sprint Plan":
    st.header("Sprint Plan")
    if not st.session_state.data: st.info("Load the synthetic workbook on Intake first.")
    else:
        n = st.number_input("Number of planned sprints", 1, 20, 7)
        if st.button("Generate dependency-safe plan", type="primary"):
            try: st.session_state.plan = plan_sprints(st.session_state.data, n)
            except Exception as e: st.error(str(e))
        if st.session_state.plan:
            plan, unresolved, capacities = st.session_state.plan
            st.subheader("Declared heuristic")
            st.write("Architecture-foundation-first, dependency-safe, priority-aware capacity packing. Active and completed sprint commitments are not changed.")
            st.dataframe(plan, use_container_width=True)
            st.subheader("Capacity")
            st.dataframe(capacities[["SprintID","AdjustedCapacityPts","CommittedPts","AvailablePts","HolidayCount"]], use_container_width=True)
            if not unresolved.empty:
                st.warning("Some tickets could not be placed.")
                st.dataframe(unresolved, use_container_width=True)
            st.download_button("Download sprint plan CSV", plan.to_csv(index=False), "sprint_plan.csv", "text/csv")

elif page == "5. Change Request":
    st.header("Mid-Sprint Change Request")
    if not st.session_state.data: st.info("Load the workbook first.")
    else:
        raw = st.text_area("New request")
        if st.button("Analyze change") and raw.strip():
            result = analyze_change(raw, st.session_state.data)
            st.json(result)
            st.checkbox("Human confirms the recommendation", key="change_confirm")

elif page == "6. Board Health":
    st.header("Existing Board Health")
    if not st.session_state.data: st.info("Load the workbook first.")
    else:
        if st.button("Run health checks", type="primary"):
            st.session_state.health = run_health_checks(st.session_state.data)
        if st.session_state.health:
            findings, capacities, dep = st.session_state.health
            st.metric("Findings", len(findings))
            st.dataframe(findings, use_container_width=True)
            st.download_button("Download findings CSV", findings.to_csv(index=False), "board_health.csv", "text/csv")

elif page == "7. Publish":
    st.header("Human-Controlled Publish")
    if not st.session_state.package: st.info("Generate and review a backlog first.")
    elif not st.session_state.approved: st.warning("Approve the internal backlog on the Backlog page first.")
    else:
        p = package_dict()
        st.json(p)
        confirm = st.checkbox("I understand this writes to the configured Jira sandbox")
        if st.button("Publish to Jira", type="primary", disabled=not confirm):
            key_map, results = publish_backlog(p["items"])
            st.subheader("Local-to-Jira key map")
            st.json(key_map)
            st.dataframe(pd.DataFrame(results), use_container_width=True)
